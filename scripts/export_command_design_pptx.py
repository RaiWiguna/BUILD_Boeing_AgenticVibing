from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from playwright.sync_api import sync_playwright


ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "docs" / "Proposal" / "ATC_FlowLens_PitchDeck_CommandDesign.html"
OUT_PPTX = ROOT / "docs" / "Proposal" / "ATC_FlowLens_PitchDeck_CommandDesign.pptx"
SLIDE_DIR = ROOT / "outputs" / "command_design_slides"


def main():
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    slide_paths = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
        page.goto(HTML.resolve().as_uri(), wait_until="networkidle")
        slides = page.locator(".slide")
        for i in range(slides.count()):
            path = SLIDE_DIR / f"slide_{i + 1:02d}.png"
            slides.nth(i).screenshot(path=str(path))
            slide_paths.append(path)
        browser.close()

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in slide_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT_PPTX)
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
