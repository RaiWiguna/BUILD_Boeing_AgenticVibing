from pathlib import Path
from math import cos, sin, pi

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "Proposal"
PPTX_OUT = OUT_DIR / "ATC_FlowLens_PitchDeck_Agentic_Vibing.pptx"
PDF_OUT = OUT_DIR / "ATC_FlowLens_PitchDeck_Agentic_Vibing.pdf"

W, H = 13.333, 7.5

COLORS = {
    "navy": "061A3A",
    "blue": "0033A0",
    "blue2": "0B5CAB",
    "sky": "D7E8FF",
    "cyan": "42A5FF",
    "ice": "F5FAFF",
    "white": "FFFFFF",
    "muted": "5B6B84",
    "line": "C8D8EC",
    "green": "1E9E6E",
    "amber": "F2A900",
    "red": "D94A3A",
}


def rgb(hex_value):
    h = hex_value.replace("#", "")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def ppt_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def ppt_line(shape, color, width=1, transparency=0):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_text(slide, text, x, y, w, h, size=18, color=None, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color or COLORS["navy"])
    return box


def add_multiline(slide, lines, x, y, w, h, size=16, color=None, bullet=False, gap=0.93):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color or COLORS["navy"])
        p.space_after = Pt(5 * gap)
        if bullet:
            p.level = 0
            p.text = f"• {line}"
    return box


def add_bg(slide, variant="light"):
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    ppt_fill(bg, COLORS["ice"] if variant == "light" else COLORS["navy"])
    bg.line.fill.background()
    # restrained technical grid
    for i in range(0, 15):
        x = i * W / 14
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), 0, Inches(x), Inches(H))
        ppt_line(line, COLORS["line"] if variant == "light" else "14365F", 0.35, 35)
    for j in range(0, 9):
        y = j * H / 8
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, Inches(y), Inches(W), Inches(y))
        ppt_line(line, COLORS["line"] if variant == "light" else "14365F", 0.35, 35)


def add_footer(slide, num, source=None, dark=False):
    c = COLORS["sky"] if dark else COLORS["muted"]
    add_text(slide, "Agentic Vibing  |  ATC FlowLens", 0.55, 7.08, 4.4, 0.2, 8.5, c)
    if source:
        add_text(slide, source, 4.7, 7.08, 6.5, 0.2, 7.5, c, align=PP_ALIGN.CENTER)
    add_text(slide, f"{num:02d}/10", 12.1, 7.08, 0.65, 0.2, 8.5, c, align=PP_ALIGN.RIGHT)


def add_kicker(slide, label, dark=False):
    color = COLORS["sky"] if dark else COLORS["blue"]
    dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.58), Inches(0.43), Inches(0.12), Inches(0.12))
    ppt_fill(dot, COLORS["cyan"])
    dot.line.fill.background()
    add_text(slide, label.upper(), 0.78, 0.36, 5.5, 0.28, 9.5, color, bold=True)


def add_title(slide, kicker, title, subtitle=None, dark=False):
    add_kicker(slide, kicker, dark)
    add_text(slide, title, 0.55, 0.82, 8.8, 0.85, 28, COLORS["white"] if dark else COLORS["navy"], bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.62, 8.5, 0.45, 13.5, COLORS["sky"] if dark else COLORS["muted"])


def card(slide, x, y, w, h, title, body, accent=COLORS["blue"], fill=COLORS["white"], title_size=15, body_size=10.5):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    ppt_fill(s, fill)
    ppt_line(s, COLORS["line"], 0.8)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    ppt_fill(bar, accent)
    bar.line.fill.background()
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.35, 0.28, title_size, COLORS["navy"], True)
    add_text(slide, body, x + 0.22, y + 0.58, w - 0.35, h - 0.72, body_size, COLORS["muted"])
    return s


def pill(slide, text, x, y, w, color=COLORS["blue"], text_color=COLORS["white"]):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36))
    ppt_fill(s, color)
    s.line.fill.background()
    add_text(slide, text, x + 0.08, y + 0.085, w - 0.16, 0.16, 9.5, text_color, True, PP_ALIGN.CENTER)


def draw_runway(slide, x, y, w, h, dark=False):
    base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    ppt_fill(base, "102D57" if dark else "EAF4FF")
    ppt_line(base, "275B93" if dark else COLORS["line"], 1.0)
    runway = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + w * 0.43), Inches(y + 0.2), Inches(w * 0.14), Inches(h - 0.4))
    ppt_fill(runway, "DDEBFA" if dark else "193B67")
    runway.line.fill.background()
    for k in range(8):
        yy = y + 0.45 + k * (h - 0.9) / 7
        mark = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + w * 0.492), Inches(yy), Inches(w * 0.018), Inches(0.18))
        ppt_fill(mark, COLORS["white"] if not dark else COLORS["navy"])
        mark.line.fill.background()
    # taxiway and tracks
    for off in [0.19, 0.81]:
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w * off), Inches(y + h * 0.15), Inches(x + w * 0.5), Inches(y + h * 0.5))
        ppt_line(line, COLORS["cyan"], 1.5, 5)
    # aircraft and vehicle
    ac = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, Inches(x + w * 0.47), Inches(y + h * 0.18), Inches(0.36), Inches(0.36))
    ppt_fill(ac, COLORS["cyan"])
    ac.rotation = 180
    ac.line.fill.background()
    veh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + w * 0.59), Inches(y + h * 0.53), Inches(0.34), Inches(0.2))
    ppt_fill(veh, COLORS["amber"])
    veh.line.fill.background()


def connector(slide, x1, y1, x2, y2, color=COLORS["blue"], width=1.4):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ppt_line(line, color, width)
    return line


def slide1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "dark")
    # diagonal light sweep
    sweep = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(7.25), Inches(-0.3), Inches(6.6), Inches(8.2))
    ppt_fill(sweep, "0B5CAB", 8)
    sweep.line.fill.background()
    draw_runway(slide, 8.35, 1.1, 3.7, 5.1, dark=True)
    pill(slide, "BUILD Indonesia  |  Aerospace & Defense", 0.65, 0.56, 3.25, COLORS["white"], COLORS["blue"])
    add_text(slide, "ATC FlowLens", 0.65, 1.65, 6.7, 0.85, 46, COLORS["white"], True)
    add_text(slide, "Explainable Runway Safety and ATC Decision-Support Dashboard", 0.68, 2.52, 5.9, 0.9, 21, COLORS["sky"], False)
    add_text(slide, "A proposal deck for BUILD Indonesia Hackathon", 0.68, 3.62, 5.1, 0.35, 14, COLORS["white"])
    add_text(slide, "Team", 0.68, 5.35, 1.0, 0.2, 10, COLORS["sky"], True)
    add_text(slide, "Agentic Vibing", 0.68, 5.65, 3.8, 0.38, 23, COLORS["white"], True)
    add_footer(slide, 1, dark=True)


def slide2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Problem", "Runway operations are time-critical", "Controllers must connect aircraft, vehicles, runway status, weather, and communication within seconds.")
    # central controller node
    center = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5.65), Inches(3.08), Inches(1.55), Inches(1.55))
    ppt_fill(center, COLORS["blue"])
    center.line.fill.background()
    add_text(slide, "ATC\nController", 5.83, 3.43, 1.18, 0.48, 16, COLORS["white"], True, PP_ALIGN.CENTER)
    nodes = [
        ("Aircraft\ntracks", 2.0, 2.0),
        ("Runway\nstatus", 4.55, 1.75),
        ("Vehicle\nmovement", 7.7, 1.95),
        ("Weather &\nvisibility", 9.72, 3.05),
        ("Radio &\nreadbacks", 7.65, 5.15),
        ("Flight\nstrips", 4.42, 5.35),
        ("Unit\ncoordination", 1.85, 4.4),
    ]
    for label, x, y in nodes:
        s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.45), Inches(0.62))
        ppt_fill(s, COLORS["white"])
        ppt_line(s, COLORS["line"])
        add_text(slide, label, x + 0.08, y + 0.1, 1.28, 0.36, 10.5, COLORS["navy"], True, PP_ALIGN.CENTER)
        connector(slide, x + 0.72, y + 0.31, 6.42, 3.86, COLORS["cyan"], 1.0)
    card(slide, 0.65, 6.08, 12.05, 0.62, "Core issue", "When information is scattered across displays, voice communication, and human memory, runway conflict risk becomes harder to detect early.", COLORS["blue"], COLORS["white"], 13, 11.5)
    add_footer(slide, 2)


def slide3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Evidence", "Runway conflict is a real safety issue", "Real incidents show how timing, communication, visibility, workload, and vehicle movement can combine into risk.")
    cases = [
        ("LATAM Peru 2213", "Lima, 2022", "Aircraft collided with a fire truck on the runway during takeoff roll.", "Aircraft-vehicle conflict near an active runway.", COLORS["red"]),
        ("Austin-Bergstrom", "2023", "Departure aircraft was cleared while another aircraft was on final approach in dense fog.", "Low visibility and runway status awareness.", COLORS["amber"]),
        ("LaGuardia", "2026 preliminary", "Reports described fire truck crossing close to a landing aircraft in a high-workload environment.", "Time-to-conflict and workload awareness.", COLORS["blue"]),
    ]
    x = 0.76
    for i, (name, date, what, relevance, color) in enumerate(cases):
        cx = x + i * 4.15
        card(slide, cx, 2.08, 3.65, 3.25, name, f"{date}\n\n{what}\n\nRelevance: {relevance}", color, COLORS["white"], 15, 10.2)
    add_text(slide, "Use as safety context, not as a claim that ATC FlowLens would have prevented these events.", 0.78, 5.75, 11.2, 0.34, 12, COLORS["muted"])
    # conflict window line
    connector(slide, 1.0, 6.45, 11.9, 6.45, COLORS["line"], 2)
    pill(slide, "Aircraft on final / takeoff", 1.0, 6.24, 2.1, COLORS["blue"])
    pill(slide, "Vehicle entering / crossing", 5.05, 6.24, 2.25, COLORS["amber"], COLORS["navy"])
    pill(slide, "Conflict window", 9.2, 6.24, 1.75, COLORS["red"])
    add_footer(slide, 3, "Sources: NTSB, Aviation Safety Network, AP preliminary reporting")


def slide4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Gap", "Current systems help, but the picture can still be fragmented", "ATC FlowLens is positioned as a support and training layer, not a replacement for certified systems.")
    add_text(slide, "Existing ATC / ATM systems", 0.75, 2.0, 4.4, 0.32, 18, COLORS["navy"], True)
    techs = ["Radar / ADS-B / MLAT", "Surface surveillance", "Electronic flight strips", "A-SMGCS / RIMCAS / ASDE-X", "STCA / MTCD", "CPDLC / Data Comm", "SWIM / flight data"]
    for i, t in enumerate(techs):
        pill(slide, t, 0.78 + (i % 2) * 2.05, 2.55 + (i // 2) * 0.58, 1.85, COLORS["sky"], COLORS["blue"])
    add_text(slide, "Remaining operational gap", 6.05, 2.0, 4.4, 0.32, 18, COLORS["navy"], True)
    gaps = [
        "Not every airport has the same surface surveillance maturity.",
        "Voice remains common in tower and ground operations.",
        "Runway status, clearance, timing, and vehicle movement may sit in separate places.",
        "Alerts are more useful when tied to an explainable event timeline.",
    ]
    add_multiline(slide, gaps, 6.08, 2.48, 5.7, 2.35, 12.2, COLORS["navy"], bullet=True)
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(5.65), Inches(9.35), Inches(0.82))
    ppt_fill(s, COLORS["blue"])
    s.line.fill.background()
    add_text(slide, "ATC FlowLens makes runway risk easier to understand through simulation, support views, and training replay.", 2.25, 5.88, 8.8, 0.26, 15, COLORS["white"], True, PP_ALIGN.CENTER)
    add_footer(slide, 4, "Technology references: FAA NextGen, EUROCONTROL A-SMGCS, ICAO GANP")


def slide5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Solution", "ATC FlowLens turns runway complexity into one operational view", "A decision-support dashboard for conflict risk, workload awareness, and event replay.")
    # dashboard mockup
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(2.0), Inches(12.0), Inches(4.58))
    ppt_fill(panel, COLORS["white"])
    ppt_line(panel, COLORS["line"])
    add_text(slide, "Runway Map", 0.98, 2.24, 2.1, 0.22, 10, COLORS["muted"], True)
    draw_runway(slide, 0.98, 2.58, 3.35, 3.36)
    add_text(slide, "Arrival / Runway Sequence", 4.68, 2.24, 3.1, 0.22, 10, COLORS["muted"], True)
    for i, (call, sec, state, col) in enumerate([("BAW123", "42s", "cleared to land", COLORS["red"]), ("SIA456", "260s", "continue approach", COLORS["blue"]), ("QFA789", "410s", "downwind", COLORS["green"])]):
        y = 2.74 + i * 0.75
        card(slide, 4.68, y, 3.35, 0.58, call, f"ETA touchdown {sec}  |  {state}", col, COLORS["ice"], 11.5, 8.5)
    add_text(slide, "Alert + Communication Timeline", 8.48, 2.24, 3.4, 0.22, 10, COLORS["muted"], True)
    alert = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.48), Inches(2.65), Inches(3.58), Inches(1.25))
    ppt_fill(alert, "FFF1F0")
    ppt_line(alert, "F0A3A0")
    add_text(slide, "HIGH RUNWAY CONFLICT", 8.68, 2.82, 3.05, 0.18, 11, COLORS["red"], True)
    add_text(slide, "Vehicle clear time exceeds aircraft touchdown time by 16 seconds.", 8.68, 3.14, 3.0, 0.46, 10, COLORS["navy"])
    timeline = [("18:42:11", "BAW123 cleared to land RWY 27"), ("18:42:18", "ARFF2 cleared to cross RWY 27"), ("18:42:21", "Alert triggered")]
    for i, (tm, msg) in enumerate(timeline):
        y = 4.2 + i * 0.47
        add_text(slide, tm, 8.52, y, 0.72, 0.16, 8.5, COLORS["blue"], True)
        add_text(slide, msg, 9.24, y, 2.75, 0.18, 8.5, COLORS["navy"])
    add_footer(slide, 5)


def slide6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "How it works", "From simulated operational data to explainable runway risk", "The MVP can prove the core logic without sensitive live ATC data.")
    cols = [
        ("Inputs", ["Aircraft position, speed, phase", "Vehicle position and status", "Runway and taxiway layout", "Clearance and communication log", "Weather and visibility"], COLORS["blue"]),
        ("Processing", ["Runway occupancy model", "Clearance state engine", "Time-to-conflict calculation", "Workload scoring", "Severity and explanation logic"], COLORS["cyan"]),
        ("Outputs", ["Runway map dashboard", "Conflict alert and reason", "Recommended action category", "Workload indicator", "Replay timeline"], COLORS["green"]),
    ]
    for i, (title, items, color) in enumerate(cols):
        x = 0.8 + i * 4.08
        card(slide, x, 2.02, 3.28, 2.78, title, "\n".join(items), color, COLORS["white"], 16, 10.8)
        if i < 2:
            connector(slide, x + 3.35, 3.36, x + 4.0, 3.36, COLORS["blue"], 2)
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.28), Inches(5.34), Inches(10.75), Inches(1.0))
    ppt_fill(rule, COLORS["navy"])
    rule.line.fill.background()
    add_text(slide, 'Example rule: IF aircraft is cleared to land AND vehicle is cleared to cross the same runway AND vehicle clear time > aircraft touchdown time, THEN trigger HIGH conflict alert.', 1.55, 5.6, 10.15, 0.36, 13, COLORS["white"], True, PP_ALIGN.CENTER)
    add_footer(slide, 6)


def slide7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "MVP", "A focused demo scenario keeps the prototype buildable", "Simulation-first scope makes the idea credible for a hackathon and safe from live-data dependency.")
    steps = [
        ("1", "Aircraft BAW123", "Final approach and cleared to land."),
        ("2", "Vehicle ARFF-2", "Receives runway crossing clearance."),
        ("3", "Prediction", "System estimates touchdown and clear time."),
        ("4", "Conflict", "Vehicle may still occupy runway."),
        ("5", "FlowLens", "Alert, explanation, workload, replay."),
    ]
    x0 = 0.72
    for i, (n, title, body) in enumerate(steps):
        x = x0 + i * 2.47
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(2.3), Inches(0.46), Inches(0.46))
        ppt_fill(circ, COLORS["blue"] if i < 3 else COLORS["red"] if i == 3 else COLORS["green"])
        circ.line.fill.background()
        add_text(slide, n, x + 0.11, 2.405, 0.24, 0.15, 10, COLORS["white"], True, PP_ALIGN.CENTER)
        if i < 4:
            connector(slide, x + 0.5, 2.53, x + 2.25, 2.53, COLORS["line"], 2)
        card(slide, x, 3.05, 2.05, 1.26, title, body, COLORS["blue"], COLORS["white"], 13, 9.2)
    add_text(slide, "Build priority", 0.82, 5.05, 2.0, 0.28, 17, COLORS["navy"], True)
    priorities = ["data model", "map visualization", "clearance engine", "time-to-conflict", "alert explanation", "workload score", "replay mode"]
    for i, p in enumerate(priorities):
        pill(slide, p, 0.85 + (i % 4) * 2.35, 5.55 + (i // 4) * 0.52, 2.05, COLORS["sky"], COLORS["blue"])
    add_footer(slide, 7)


def slide8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Users & business impact", "Start with training and safety awareness before live operations", "A realistic adoption path creates value while avoiding premature certification claims.")
    users = ["Tower controllers", "Ground controllers", "Approach controllers", "Airport operation centers", "ARFF units", "ATC training centers", "Safety investigators"]
    for i, u in enumerate(users):
        pill(slide, u, 0.85 + (i % 2) * 2.65, 2.05 + (i // 2) * 0.56, 2.38, COLORS["white"], COLORS["blue"])
    add_text(slide, "Initial market positioning", 6.55, 2.05, 3.8, 0.3, 17, COLORS["navy"], True)
    card(slide, 6.55, 2.48, 5.38, 1.05, "Simulation, training, and runway safety awareness", "For airports, aviation schools, ANSPs, and airport operation teams.", COLORS["blue"], COLORS["white"], 15, 11)
    values = [
        ("Training quality", "Better runway conflict scenarios."),
        ("Review speed", "Clearer event timelines."),
        ("Operational readiness", "Safer tower-ground-emergency coordination."),
        ("Expansion path", "Speech-to-text, readback mismatch, richer analytics."),
    ]
    for i, (t, b) in enumerate(values):
        card(slide, 0.85 + (i % 2) * 5.45, 5.0 + (i // 2) * 0.78, 4.85, 0.58, t, b, COLORS["cyan"], COLORS["ice"], 12, 9.2)
    add_footer(slide, 8)


def slide9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Impact & alignment", "Advanced technology applied to a real aerospace safety challenge", "The concept connects safety, training, operational readiness, and sustainability.")
    impacts = [
        ("Earlier risk recognition", "See aircraft, vehicle, clearance, and timing in one view.", COLORS["blue"]),
        ("Better workload awareness", "Expose active risk factors before task saturation.", COLORS["cyan"]),
        ("Stronger training replay", "Review why a conflict happened and what contributed.", COLORS["green"]),
        ("Sustainability upside", "Future versions may reduce avoidable go-arounds, taxi delay, holding, and fuel burn.", COLORS["amber"]),
    ]
    for i, (t, b, c) in enumerate(impacts):
        card(slide, 0.75 + (i % 2) * 5.9, 2.0 + (i // 2) * 1.28, 5.25, 0.95, t, b, c, COLORS["white"], 14, 10.2)
    align = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.02), Inches(5.25), Inches(11.25), Inches(0.95))
    ppt_fill(align, COLORS["blue"])
    align.line.fill.background()
    add_text(slide, "BUILD Indonesia alignment", 1.35, 5.43, 2.5, 0.2, 12, COLORS["sky"], True)
    add_text(slide, "Aerospace & Defense  |  Advanced Technology and Innovation  |  Social Impact and Sustainability", 1.35, 5.75, 9.9, 0.24, 15, COLORS["white"], True)
    add_footer(slide, 9)


def slide10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "dark")
    add_title(slide, "Roadmap", "From hackathon MVP to aviation safety training platform", dark=True)
    phases = [
        ("Phase 1", "Hackathon MVP", ["simulated runway map", "aircraft and vehicle movement", "time-to-conflict alert", "workload indicator", "replay timeline"], COLORS["cyan"]),
        ("Phase 2", "Enhanced prototype", ["ATC-style speech-to-text", "readback mismatch detection", "scenario generator", "improved workload model", "multi-runway support"], COLORS["sky"]),
        ("Phase 3", "Future vision", ["airport operation datasets", "tower/ground training simulator", "safety analytics dashboard", "research path to certified support"], COLORS["green"]),
    ]
    for i, (phase, title, items, color) in enumerate(phases):
        x = 0.75 + i * 4.1
        s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.1), Inches(3.35), Inches(2.95))
        ppt_fill(s, "0D2A55")
        ppt_line(s, "214C7C", 1)
        add_text(slide, phase.upper(), x + 0.25, 2.35, 1.3, 0.2, 9.5, color, True)
        add_text(slide, title, x + 0.25, 2.72, 2.75, 0.28, 17, COLORS["white"], True)
        add_multiline(slide, items, x + 0.25, 3.2, 2.8, 1.35, 10, COLORS["sky"], bullet=True)
    add_text(slide, "ATC FlowLens helps air traffic teams see runway risk before it becomes a runway incident.", 0.8, 5.86, 11.6, 0.42, 22, COLORS["white"], True, PP_ALIGN.CENTER)
    add_text(slide, "Agentic Vibing", 0.8, 6.58, 11.6, 0.22, 12, COLORS["sky"], True, PP_ALIGN.CENTER)
    add_footer(slide, 10, dark=True)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for fn in [slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10]:
        fn(prs)
    prs.save(PPTX_OUT)


# PDF helpers use points.
PW, PH = landscape((13.333 * inch, 7.5 * inch))


def pdf_text(c, text, x, y, size=12, color=COLORS["navy"], bold=False, align="left"):
    c.setFillColor(colors.HexColor("#" + color))
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    tx = x * inch
    ty = PH - y * inch
    if align == "center":
        c.drawCentredString(tx, ty, text)
    elif align == "right":
        c.drawRightString(tx, ty, text)
    else:
        c.drawString(tx, ty, text)


def pdf_wrapped(c, text, x, y, w, size=10, color=COLORS["navy"], bold=False, leading=1.22):
    words = text.split()
    lines, line = [], ""
    max_chars = max(18, int(w * 10.5 * 12 / size))
    for word in words:
        candidate = (line + " " + word).strip()
        if len(candidate) > max_chars and line:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    for i, line in enumerate(lines):
        pdf_text(c, line, x, y + i * size * leading / 72, size, color, bold)


def pdf_rect(c, x, y, w, h, fill, stroke=None, radius=0):
    c.setFillColor(colors.HexColor("#" + fill))
    c.setStrokeColor(colors.HexColor("#" + (stroke or fill)))
    if radius:
        c.roundRect(x * inch, PH - (y + h) * inch, w * inch, h * inch, radius * inch, fill=1, stroke=1 if stroke else 0)
    else:
        c.rect(x * inch, PH - (y + h) * inch, w * inch, h * inch, fill=1, stroke=1 if stroke else 0)


def pdf_bg(c, dark=False):
    pdf_rect(c, 0, 0, W, H, COLORS["navy"] if dark else COLORS["ice"])
    c.setStrokeColor(colors.HexColor("#214C7C" if dark else "#" + COLORS["line"]))
    c.setLineWidth(0.25)
    for i in range(15):
        x = i * W / 14 * inch
        c.line(x, 0, x, PH)
    for j in range(9):
        y = j * H / 8 * inch
        c.line(0, y, PW, y)


def pdf_title(c, kicker, title, subtitle=None, dark=False):
    pdf_text(c, kicker.upper(), 0.78, 0.55, 9.5, COLORS["sky"] if dark else COLORS["blue"], True)
    pdf_text(c, title, 0.58, 1.18, 28, COLORS["white"] if dark else COLORS["navy"], True)
    if subtitle:
        pdf_wrapped(c, subtitle, 0.6, 1.68, 8.5, 12.5, COLORS["sky"] if dark else COLORS["muted"])


def pdf_footer(c, num, dark=False):
    col = COLORS["sky"] if dark else COLORS["muted"]
    pdf_text(c, "Agentic Vibing  |  ATC FlowLens", 0.55, 7.22, 8.5, col)
    pdf_text(c, f"{num:02d}/10", 12.65, 7.22, 8.5, col, align="right")


def pdf_card(c, x, y, w, h, title, body, accent=COLORS["blue"], fill=COLORS["white"]):
    pdf_rect(c, x, y, w, h, fill, COLORS["line"], 0.08)
    pdf_rect(c, x, y, 0.08, h, accent)
    pdf_wrapped(c, title, x + 0.2, y + 0.3, w - 0.35, 12, COLORS["navy"], True)
    pdf_wrapped(c, body, x + 0.2, y + 0.72, w - 0.35, 9.2, COLORS["muted"])


def build_pdf():
    c = canvas.Canvas(str(PDF_OUT), pagesize=(PW, PH))

    # PDF pages mirror the PPTX but use simplified vector layouts.
    pdf_bg(c, True)
    pdf_rect(c, 7.8, 0, 5.6, 7.5, COLORS["blue2"])
    pdf_text(c, "BUILD Indonesia  |  Aerospace & Defense", 0.7, 0.8, 11, COLORS["sky"], True)
    pdf_text(c, "ATC FlowLens", 0.7, 2.18, 43, COLORS["white"], True)
    pdf_wrapped(c, "Explainable Runway Safety and ATC Decision-Support Dashboard", 0.72, 2.75, 5.4, 20, COLORS["sky"])
    pdf_text(c, "Team", 0.72, 5.45, 10, COLORS["sky"], True)
    pdf_text(c, "Agentic Vibing", 0.72, 5.9, 24, COLORS["white"], True)
    pdf_footer(c, 1, True); c.showPage()

    pages = [
        ("Problem", "Runway operations are time-critical", "Controllers must connect aircraft, vehicles, runway status, weather, and communication within seconds."),
        ("Evidence", "Runway conflict is a real safety issue", "Real incidents show how timing, communication, visibility, workload, and vehicle movement can combine into risk."),
        ("Gap", "Current systems help, but the picture can still be fragmented", "ATC FlowLens is a support and training layer, not a replacement for certified systems."),
        ("Solution", "ATC FlowLens turns runway complexity into one operational view", "A dashboard for conflict risk, workload awareness, and event replay."),
        ("How it works", "From simulated operational data to explainable runway risk", "The MVP can prove the core logic without sensitive live ATC data."),
        ("MVP", "A focused demo scenario keeps the prototype buildable", "Simulation-first scope makes the idea credible for a hackathon."),
        ("Users & business impact", "Start with training and safety awareness before live operations", "A realistic adoption path creates value while avoiding premature certification claims."),
        ("Impact & alignment", "Advanced technology applied to a real aerospace safety challenge", "The concept connects safety, training, operational readiness, and sustainability."),
    ]
    page_bodies = [
        [
            ("Aircraft tracks", "Approaching, landing, departing, or taxiing aircraft."),
            ("Runway status", "Active runway, occupancy, crossing, takeoff, and landing state."),
            ("Vehicle movement", "Airport vehicles and runway crossing requests."),
            ("Weather and visibility", "Fog, wind shift, storm cells, and low-visibility operations."),
            ("Radio and readbacks", "Voice instructions, blocked transmissions, and readback risk."),
            ("Flight strips", "Clearance, runway, sequence, and handoff information."),
            ("Unit coordination", "Ground, tower, approach, airport ops, and emergency teams."),
        ],
        [
            ("LATAM Peru 2213", "Lima, 2022: aircraft collided with a fire truck on runway during takeoff roll."),
            ("Austin-Bergstrom", "2023: departure aircraft was cleared while another aircraft was on final in dense fog."),
            ("LaGuardia 2026 preliminary", "Reports described fire truck crossing close to landing aircraft in high workload."),
            ("Common pattern", "Timing pressure, communication, vehicle movement, visibility, and workload combine."),
        ],
        [
            ("Existing systems", "Radar, ADS-B, MLAT, surface surveillance, electronic strips, A-SMGCS, RIMCAS, STCA, CPDLC, SWIM."),
            ("Maturity gap", "Not every airport has the same surface surveillance and conflict alerting capability."),
            ("Voice remains common", "Tower and ground operations still depend heavily on radio communication."),
            ("Fragmentation", "Runway status, clearance, aircraft timing, and vehicle movement can sit in separate places."),
            ("FlowLens role", "Support, simulation, and training layer that makes runway risk easier to understand."),
        ],
        [
            ("Runway occupancy", "Shows clear, occupied, crossing, takeoff, or inbound landing state."),
            ("Aircraft and vehicle tracking", "Displays simulated positions on a simple airport map."),
            ("Clearance state engine", "Tracks cleared to land, cleared for takeoff, hold short, cross, and go around."),
            ("Time-to-conflict", "Compares aircraft touchdown or runway entry time with vehicle clear time."),
            ("Explainable alert", "Shows why an alert appears, not only that risk exists."),
            ("Replay mode", "Reviews the event timeline for training and safety analysis."),
        ],
        [
            ("Inputs", "Aircraft position, speed, phase, runway; vehicle position, status, route; clearance, comms, weather."),
            ("Processing", "Runway occupancy model, clearance engine, time-to-conflict, workload score, severity classification."),
            ("Outputs", "Dashboard, conflict alert, reason, recommended action category, workload indicator, replay timeline."),
            ("Example rule", "If cleared-to-land aircraft arrives before crossed vehicle clears the same runway, trigger HIGH alert."),
        ],
        [
            ("BAW123", "Aircraft is on final approach and cleared to land."),
            ("ARFF-2", "Airport vehicle receives runway crossing clearance."),
            ("Prediction", "System estimates aircraft touchdown time and vehicle runway clear time."),
            ("Conflict", "Vehicle may still occupy runway when aircraft arrives."),
            ("FlowLens output", "Alert, explanation, workload state, and replay timeline."),
            ("Build priority", "Data model, map, clearance engine, conflict detection, alert explanation, workload, replay."),
        ],
        [
            ("Target users", "Tower, ground, approach, airport ops, ARFF, ATC training centers, safety investigators."),
            ("Initial market", "Simulation, training, and runway safety awareness for airports, aviation schools, ANSPs, and airport teams."),
            ("Training value", "Improves runway conflict scenario quality and post-event explanation."),
            ("Operational value", "Supports safety briefings, after-action review, and tower-ground-emergency coordination."),
            ("Expansion path", "Speech-to-text, readback mismatch detection, richer workload analytics."),
        ],
        [
            ("Earlier risk recognition", "See aircraft movement, vehicle activity, clearance status, and timing in one view."),
            ("Better workload awareness", "Expose active risk factors before task saturation."),
            ("Stronger training replay", "Review why a conflict happened and what contributed."),
            ("BUILD alignment", "Aerospace & Defense, Advanced Technology and Innovation, Social Impact, Sustainability."),
            ("Future sustainability", "Potential to reduce avoidable go-arounds, taxi delay, holding, and fuel burn."),
        ],
    ]
    for idx, (k, t, s) in enumerate(pages, start=2):
        pdf_bg(c)
        pdf_title(c, k, t, s)
        for i, item in enumerate(page_bodies[idx - 2]):
            x = 0.95 + (i % 2) * 5.55
            y = 2.35 + (i // 2) * 0.9
            title, body = item
            pdf_card(c, x, y, 4.9, 0.72, title, body, COLORS["blue"] if i % 3 == 0 else COLORS["cyan"])
        pdf_footer(c, idx)
        c.showPage()

    pdf_bg(c, True)
    pdf_title(c, "Roadmap", "From hackathon MVP to aviation safety training platform", dark=True)
    phase_text = [
        ("Phase 1", "Hackathon MVP: simulated runway map, time-to-conflict alert, workload indicator, replay timeline."),
        ("Phase 2", "Enhanced prototype: speech-to-text, readback mismatch, scenario generator, multi-runway support."),
        ("Phase 3", "Future vision: airport data, tower-ground training simulator, safety analytics, certified support research."),
    ]
    for i, (title, body) in enumerate(phase_text):
        pdf_card(c, 0.9 + i * 4.1, 2.45, 3.35, 2.2, title, body, COLORS["cyan"], "0D2A55")
    pdf_text(c, "ATC FlowLens helps air traffic teams see runway risk before it becomes a runway incident.", 6.66, 6.1, 21, COLORS["white"], True, "center")
    pdf_text(c, "Agentic Vibing", 6.66, 6.72, 12, COLORS["sky"], True, "center")
    pdf_footer(c, 10, True); c.showPage()
    c.save()


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_pptx()
    build_pdf()
    print(PPTX_OUT)
    print(PDF_OUT)
